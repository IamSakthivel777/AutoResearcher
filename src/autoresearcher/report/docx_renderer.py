"""Deterministic rendering of validated report specifications."""

from __future__ import annotations

from contextlib import suppress
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from autoresearcher.report.schema import (
    BulletsBlock,
    CalloutBlock,
    ChecklistBlock,
    DayPlanBlock,
    HeadingBlock,
    ImageBlock,
    ImageGalleryBlock,
    NumberedBlock,
    PageBreakBlock,
    ParagraphBlock,
    ReportSpec,
    TableBlock,
)

if TYPE_CHECKING:
    from docx.text.paragraph import Paragraph


class DocxRenderer:
    """Render ``ReportSpec`` to a Word-compatible document."""

    heading_color = RGBColor(31, 78, 121)
    accent_fill = "D9EAF7"

    def render(self, report: ReportSpec, output: str | Path) -> Path:
        """Write ``report`` to ``output`` and return the resolved file path."""
        path = Path(output).expanduser().resolve()
        path.parent.mkdir(parents=True, exist_ok=True)

        document = Document()
        self._configure(document)
        self._title_page(document, report)
        self._table_of_contents(document)

        source_numbers = {source.id: index for index, source in enumerate(report.sources, 1)}
        for section in report.sections:
            heading = document.add_heading(section.title, level=1)
            self._apply_language_font(heading)
            for block in section.blocks:
                self._render_block(document, block, source_numbers)

        self._sources(document, report)
        self._header_footer(document, report.title)
        document.save(str(path))
        return path

    def _configure(self, document: DocumentObject) -> None:
        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor(45, 55, 65)
        normal.paragraph_format.space_after = Pt(6)
        for level in range(1, 5):
            style = styles[f"Heading {level}"]
            style.font.name = "Aptos Display"
            style.font.color.rgb = self.heading_color
            style.font.bold = True
            if level == 1:
                self._bottom_border(style.element, "2B7A78")
        if "Source Marker" not in styles:
            marker = styles.add_style("Source Marker", WD_STYLE_TYPE.CHARACTER)
            marker.font.size = Pt(8)
            marker.font.color.rgb = self.heading_color

        section = document.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    def _title_page(self, document: DocumentObject, report: ReportSpec) -> None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(48 if report.cover_image else 140)
        title = paragraph.add_run(report.title)
        title.bold = True
        title.font.size = Pt(28)
        title.font.color.rgb = self.heading_color
        if report.subtitle:
            subtitle = document.add_paragraph(report.subtitle)
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.runs[0].font.size = Pt(15)
        date_paragraph = document.add_paragraph(
            f"Prepared on {report.report_date.strftime('%d %B %Y')}"
        )
        date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if report.cover_image is not None:
            self._render_image(document, report.cover_image, {
                source.id: index for index, source in enumerate(report.sources, 1)
            })
        if report.assumptions:
            document.add_heading("Assumptions", level=2)
            for assumption in report.assumptions:
                document.add_paragraph(assumption, style="List Bullet")
        document.add_page_break()  # type: ignore[no-untyped-call]

    def _table_of_contents(self, document: DocumentObject) -> None:
        document.add_heading("Table of Contents", level=1)
        paragraph = document.add_paragraph()
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        placeholder = OxmlElement("w:t")
        placeholder.text = "Right-click and select Update Field."
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for element in (begin, instruction, separate, placeholder, end):
            run._r.append(element)
        document.add_page_break()  # type: ignore[no-untyped-call]

    def _render_block(
        self, document: DocumentObject, block: object, source_numbers: dict[str, int]
    ) -> None:
        if isinstance(block, HeadingBlock):
            paragraph = document.add_heading(block.text, level=block.level)
            self._apply_language_font(paragraph)
        elif isinstance(block, ParagraphBlock):
            paragraph = document.add_paragraph(block.text)
            self._markers(paragraph, block.source_ids, source_numbers)
            self._apply_language_font(paragraph)
        elif isinstance(block, (BulletsBlock, NumberedBlock, ChecklistBlock)):
            style = "List Bullet" if isinstance(block, BulletsBlock) else "List Number"
            for item in block.items:
                text = f"☐ {item}" if isinstance(block, ChecklistBlock) else item
                paragraph = document.add_paragraph(text, style=style)
                self._markers(paragraph, block.source_ids, source_numbers)
                self._apply_language_font(paragraph)
        elif isinstance(block, TableBlock):
            table = document.add_table(rows=1, cols=len(block.headers))
            table.style = "Table Grid"
            for cell, header in zip(table.rows[0].cells, block.headers, strict=True):
                cell.text = header
                self._shade(cell._tc, self.accent_fill)
            for values in block.rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, values, strict=True):
                    cell.text = value
            if block.source_ids:
                paragraph = document.add_paragraph()
                self._markers(paragraph, block.source_ids, source_numbers)
        elif isinstance(block, CalloutBlock):
            table = document.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            cell = table.cell(0, 0)
            cell.text = f"{block.kind.upper()}: {block.text}"
            fill = {"info": "D9EAF7", "tip": "E2F0D9", "warning": "FFF2CC"}[block.kind]
            self._shade(cell._tc, fill)
            self._markers(cell.paragraphs[0], block.source_ids, source_numbers)
            self._apply_language_font(cell.paragraphs[0])
        elif isinstance(block, DayPlanBlock):
            heading = document.add_heading(f"Day {block.day}: {block.title}", level=2)
            self._apply_language_font(heading)
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            headers = ("Time", "Activity", "Details")
            for cell, header in zip(table.rows[0].cells, headers, strict=True):
                cell.text = header
                self._shade(cell._tc, self.accent_fill)
            for entry in block.entries:
                cells = table.add_row().cells
                cells[0].text = entry.time
                cells[1].text = entry.activity
                cells[2].text = entry.details
                self._markers(cells[2].paragraphs[0], entry.source_ids, source_numbers)
        elif isinstance(block, ImageBlock):
            self._render_image(document, block, source_numbers)
        elif isinstance(block, ImageGalleryBlock):
            self._render_gallery(document, block, source_numbers)
        elif isinstance(block, PageBreakBlock):
            document.add_page_break()  # type: ignore[no-untyped-call]
        else:  # pragma: no cover - the discriminated schema prevents this
            raise TypeError(f"Unsupported report block: {type(block).__name__}")

    def _markers(
        self, paragraph: Paragraph, source_ids: list[str], source_numbers: dict[str, int]
    ) -> None:
        for source_id in dict.fromkeys(source_ids):
            run = paragraph.add_run(f" [{source_numbers[source_id]}]")
            run.style = "Source Marker"

    @staticmethod
    def _apply_language_font(paragraph: Paragraph) -> None:
        """Use a Word font with coverage for the supported Indic scripts."""
        indic_ranges = (
            ("\u0900", "\u097f"),  # Devanagari
            ("\u0b80", "\u0bff"),  # Tamil
            ("\u0c00", "\u0c7f"),  # Telugu
            ("\u0d00", "\u0d7f"),  # Malayalam
        )
        for run in paragraph.runs:
            has_indic_text = any(
                start <= character <= end
                for start, end in indic_ranges
                for character in run.text
            )
            if not has_indic_text:
                continue
            run.font.name = "Nirmala UI"
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{attribute}"), "Nirmala UI")

    def _sources(self, document: DocumentObject, report: ReportSpec) -> None:
        document.add_heading("Sources", level=1)
        if not report.sources:
            document.add_paragraph("No external sources were used.")
            return
        for index, source in enumerate(report.sources, 1):
            paragraph = document.add_paragraph()
            paragraph.add_run(f"[{index}] {source.title}. ").bold = True
            self._hyperlink(paragraph, str(source.url), str(source.url))
            if source.title.startswith("Image:") and source.snippet:
                paragraph.add_run(f" — {source.snippet}")
            paragraph.add_run(f" (retrieved {source.retrieved_at.date().isoformat()})")

    def _header_footer(self, document: DocumentObject, title: str) -> None:
        for section in document.sections:
            section.header.paragraphs[0].text = title
            paragraph = section.footer.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run("Page ")
            self._field(paragraph, "PAGE")
            paragraph.add_run(" of ")
            self._field(paragraph, "NUMPAGES")

    @staticmethod
    def _field(paragraph: Paragraph, instruction: str) -> None:
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        text = OxmlElement("w:instrText")
        text.set(qn("xml:space"), "preserve")
        text.text = instruction
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend((begin, text, end))

    def _render_image(
        self,
        document: DocumentObject,
        block: ImageBlock,
        source_numbers: dict[str, int],
    ) -> None:
        path = Path(block.path)
        if not path.is_file():
            return
        try:
            document.add_picture(str(path), width=Inches(block.width_inches))
        except (UnrecognizedImageError, OSError, ValueError):
            return
        image_paragraph = document.paragraphs[-1]
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if block.area_name:
            area = caption.add_run(block.area_name)
            area.bold = True
            caption.add_run(" — ")
        run = caption.add_run(block.caption)
        run.italic = True
        run.font.size = Pt(8.5)
        self._markers(caption, block.source_ids, source_numbers)

    def _render_gallery(
        self,
        document: DocumentObject,
        block: ImageGalleryBlock,
        source_numbers: dict[str, int],
    ) -> None:
        """Render a compact two-column gallery with named-place captions."""
        rows = (len(block.images) + 1) // 2
        table = document.add_table(rows=rows, cols=2)
        table.style = "Table Grid"
        for index, image in enumerate(block.images):
            cell = table.cell(index // 2, index % 2)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            self._shade(cell._tc, "F4F8FB")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            path = Path(image.path)
            if path.is_file():
                with suppress(UnrecognizedImageError, OSError, ValueError):
                    paragraph.add_run().add_picture(str(path), width=Inches(2.85))
            caption = cell.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label = caption.add_run(image.area_name or image.alt_text)
            label.bold = True
            label.font.color.rgb = self.heading_color
            details = caption.add_run(f"\n{image.caption}")
            details.italic = True
            details.font.size = Pt(8)
            self._markers(caption, image.source_ids, source_numbers)

    @staticmethod
    def _bottom_border(style_element: object, color: str) -> None:
        properties = style_element.get_or_add_pPr()  # type: ignore[attr-defined]
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), color)
        borders.append(bottom)
        properties.append(borders)

    @staticmethod
    def _shade(table_cell: object, fill: str) -> None:
        properties = table_cell.get_or_add_tcPr()  # type: ignore[attr-defined]
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    @staticmethod
    def _hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
        relationship_id = paragraph.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.extend((color, underline))
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = text
        run.append(value)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
